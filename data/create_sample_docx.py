"""Generate a sample DOCX for testing the terminology checker."""

from docx import Document

doc = Document()
doc.add_heading("Installation Guide", 0)

doc.add_heading("1. Introduction", level=1)
doc.add_paragraph(
    "This guide describes how to set up and configure the device. "
    "Make sure you read all instructions prior to starting the installation process. "
    "Due to the fact that the installation requires administrator privileges, "
    "ensure that you have the necessary permissions."
)

doc.add_heading("2. Prerequisites", level=1)
doc.add_paragraph("In order to install the software, you must have the following:")
doc.add_paragraph("  - A compatible operating system (e.g., Windows 10 or later)")
doc.add_paragraph("  - At least 4 GB of RAM")
doc.add_paragraph("  - Administrator account")

doc.add_heading("3. Installation Steps", level=1)
doc.add_paragraph(
    "Step 1: Start up the installer by clicking on the setup file. "
    "The installer will utilise the default configuration settings."
)
doc.add_paragraph(
    "Step 2: On a regular basis, the installer checks for updates. "
    "If an update is available, it will be installed subsequent to the main installation."
)
doc.add_paragraph(
    "Step 3: After the installation is complete, set up the network connection. "
    "Make sure the device is connected to the network prior to proceeding."
)

doc.add_heading("4. Troubleshooting", level=1)
doc.add_paragraph(
    "If the installation fails, shut down the device and restart the the installer. "
    "For additional support, i.e., contacting the help desk, refer to the support section."
)

out = "data/sample_document.docx"
doc.save(out)
print(f"Created: {out}")
